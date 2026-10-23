"""Bundle every CSV in data/ into one spreadsheet with a contents sheet."""
import csv
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .config import DATA

SHEETS = [
    ("01_supply_check.csv", "Supply check", "All balances read against the 10 billion RBNT created at genesis (section 1)"),
    ("02_where_the_rbnt_is.csv", "Where RBNT sits", "RBNT by the address that physically holds it (section 2)"),
    ("03_who_holds_the_rbnt.csv", "Who holds RBNT", "Holdings with WRBNT and stakes counted with their owner (figure 1)"),
    ("04_holders_by_size.csv", "Holders by size", "Wallets and RBNT per holding size, month by month (figure 3)"),
    ("05_concentration.csv", "Concentration", "Wallet count, largest-wallet shares and Gini, month by month (figures 2, 4)"),
    ("06_cohorts.csv", "Cohorts", "The same wallets followed through the quarter (figure 5)"),
    ("06b_band_moves.csv", "Band moves", "Wallets that ended in a bigger or smaller size band (section 5)"),
    ("06c_contracts_by_kind.csv", "Contracts by kind", "Contracts grouped by the name in their verified code (section 2)"),
    ("06d_net_change_by_group.csv", "Net change", "Start, end and change for every group; changes sum to zero (figures 1, 6)"),
    ("07_weekly_named_addresses.csv", "Weekly named", "Weekly balances of reserved wallets, pools and WRBNT supply (figure 7)"),
    ("08_named_addresses.csv", "Named addresses", "Start, end, flows, fees and reconciliation of each named address (figure 10)"),
    ("09_reserved_wallet_flows.csv", "Reserved flows", "Reserved wallet transfers grouped by counterparty (section 7)"),
    ("10_staking_pools.csv", "Staking pools", "Pool balances, stakes, unstakes, rewards and top-ups (figure 8)"),
    ("11_lock_calendar.csv", "Lock calendar", "Open stakes by the month their lock ends (figure 9)"),
    ("12_group_flows.csv", "Group flows", "RBNT and WRBNT moved between groups during the quarter (section 5)"),
]

INK = "0F1A22"
STEEL = "DDE6ED"


def as_value(x):
    try:
        if x.strip() == "":
            return x
        f = float(x)
        return int(f) if f.is_integer() and "." not in x else f
    except (ValueError, AttributeError):
        return x


def build(path):
    figs = json.loads((DATA / "figures.json").read_text())
    wb = Workbook()
    ws = wb.active
    ws.title = "Contents"
    ws["A1"] = "RBNT Holder Report, Q3 2026: data"
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    ws["A3"] = (f"All figures read from Redbelly mainnet (chain ID 151) at block {figs['start_block']:,} "
                f"({figs['start_time']}) and block {figs['end_block']:,} ({figs['end_time']}), plus month-end and weekly "
                f"snapshots listed in each sheet. Data read {figs['read_on']}.")
    ws["A4"] = "RBNT amounts are whole RBNT (1 RBNT = 10^18 wei) unless a column says wei. Only publicly named addresses appear."
    ws["A6"], ws["B6"], ws["C6"] = "Sheet", "File", "What it holds"
    for c in ("A6", "B6", "C6"):
        ws[c].font = Font(bold=True)
        ws[c].fill = PatternFill("solid", fgColor=STEEL)
    for i, (fname, title, what) in enumerate(SHEETS, start=7):
        ws.cell(i, 1, title)
        ws.cell(i, 2, fname)
        ws.cell(i, 3, what)
        with open(DATA / fname) as f:
            data = list(csv.reader(f))
        s = wb.create_sheet(title)
        for r, row in enumerate(data, start=1):
            for c, v in enumerate(row, start=1):
                cell = s.cell(r, c, as_value(v) if r > 1 else v)
                if r == 1:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor=STEEL)
                    cell.alignment = Alignment(wrap_text=True, vertical="bottom")
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.##" if isinstance(cell.value, float) else "#,##0"
        for c in range(1, len(data[0]) + 1):
            s.column_dimensions[s.cell(1, c).column_letter].width = 24 if c > 1 else 40
        s.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 90
    wb.save(path)
